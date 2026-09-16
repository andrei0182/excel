"""
Macro-N — Module A/C prototype v2
Refines subject detection: a "normalized" top-level Subject (the kind that
must have a matching block in the Normalize sheet) is a row where column H
contains the marker text 'דירוג משוקלל מנורמל' (visible in the PDF sample,
e.g. Main!H174). Rows using a different calculation (e.g. direct reference
'=I90', module D/E patterns) are NOT expected to appear in Normalize — so we
exclude them, instead of naively treating every integer in column B as a
subject.
"""
import openpyxl, glob, os, difflib

MARKER = "דירוג משוקלל מנורמל"

def collect_normalized_subjects(ws):
    """Top-level subjects that use the Normalize-sheet formula pattern."""
    out = []
    for row in range(1, ws.max_row + 1):
        h = ws[f"H{row}"].value
        c = ws[f"C{row}"].value
        if isinstance(h, str) and MARKER in h and isinstance(c, str) and c.strip():
            out.append((row, c.strip()))
    return out

def collect_normalize_blocks(ws):
    """In Normalize sheet, subject header rows = integer in B + text in C,
    with no numeric min/mid/max on that same row (those live on the sub-item
    rows below it)."""
    out = []
    for row in range(1, ws.max_row + 1):
        b = ws[f"B{row}"].value
        c = ws[f"C{row}"].value
        g = ws[f"G{row}"].value
        is_int = isinstance(b, int) or (isinstance(b, float) and b.is_integer())
        if is_int and isinstance(c, str) and c.strip() and g is None:
            out.append((row, c.strip()))
    return out

def audit_file(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    main, norm = wb["Main"], wb["Normalize"]

    main_subj = collect_normalized_subjects(main)
    norm_subj = collect_normalize_blocks(norm)
    norm_texts = [t for _, t in norm_subj]

    print(f"\n{'='*70}\n{os.path.basename(path)}\n{'='*70}")
    print(f"  Normalized-type subjects in Main: {len(main_subj)}   Subject blocks in Normalize: {len(norm_subj)}")

    for row, text in main_subj:
        if text in norm_texts:
            print(f"  ✔ row {row:4d}  {text}")
        else:
            close = difflib.get_close_matches(text, norm_texts, n=1, cutoff=0.5)
            if close:
                print(f"  ✗ row {row:4d}  Main: {text!r}\n                Normalize (closest): {close[0]!r}  <-- needs sync")
            else:
                print(f"  ✗ row {row:4d}  Main: {text!r}  <-- NOT FOUND in Normalize at all")

if __name__ == "__main__":
    for f in sorted(glob.glob("/mnt/user-data/uploads/*.xlsx")):
        audit_file(f)
