#!/usr/bin/env python3
"""
Parse a Fiserv / CardPointe "Your Card Processing Statement" PDF into
structured tables (Excel, CSV, JSON) and reconcile the sections against
each other so parsing errors are caught automatically.

Usage:
    python parse_fiserv_statement.py statement.pdf                # -> statement_parsed.xlsx
    python parse_fiserv_statement.py statement.pdf -o out.xlsx --csv out_csv --json out.json
    python parse_fiserv_statement.py folder_with_pdfs/            # parse every PDF in the folder

Requires: pip install pymupdf openpyxl
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pymupdf

# --------------------------------------------------------------------------- helpers

MONEY_RE = re.compile(r"^-?\$-?[\d,]+\.\d{2}$")
DATE_RE = re.compile(r"^\d{2}/\d{2}/\d{4}$")
PCT_RE = re.compile(r"^-?[\d.]+%$")
INT_RE = re.compile(r"^-?[\d,]+$")
LOC_RE = re.compile(r"^\d{10,16}$")

SECTION_TITLES = {
    "Summary By Day": "daily",
    "Summary by Card Type": "card_type",
    "Summary By Location": "location",
    "Paid by Others": "paid_by_others",
    "Disputes": "disputes",
    "Fee Summary": "fee_summary",
    "Fees": "fees",
    "Interchange Charges / Program Fees": "interchange",
    "Amounts Funded": "amounts_funded",
}
TITLE_X = (40, 42)          # section titles sit at x≈41, table cells at x≈38, chart labels at x≥43
TOLERANCE = Decimal("0.01")


def money(s: str) -> Decimal:
    s = s.replace("$", "").replace(",", "")
    return Decimal(s)


def is_money(s: str) -> bool:
    return bool(MONEY_RE.match(s))


def to_int(s: str) -> int:
    return int(s.replace(",", ""))


def iso(d: str) -> str:
    return datetime.strptime(d, "%m/%d/%Y").date().isoformat()


def pct(s: str) -> Decimal:
    return Decimal(s.rstrip("%")) / 100


@dataclass
class Row:
    page: int
    y: float
    cells: list[tuple[int, str]]  # (x, text) sorted by x

    @property
    def texts(self) -> list[str]:
        return [t for _, t in self.cells]

    @property
    def first_x(self) -> int:
        return self.cells[0][0]


def visible_spans(page: pymupdf.Page) -> tuple[list[tuple[float, int, str]], int]:
    """
    Text spans as (top y, left x, text), keeping only characters that are actually visible.
    Characters painted over by a filled shape drawn later (e.g. white boxes used to mask a merchant
    name) are dropped; they are still in the PDF text layer, so naive extraction would leak them.
    Returns the spans and the number of hidden characters.
    """
    covers = [(d["seqno"], d["rect"]) for d in page.get_drawings() if d.get("fill") is not None]
    spans, hidden = [], 0
    for s in page.get_texttrace():
        later = [r for seq, r in covers if seq > s["seqno"] and r.intersects(s["bbox"])]
        chars = []
        for ch in s["chars"]:
            box = pymupdf.Rect(ch[3])
            c = pymupdf.Point((box.x0 + box.x1) / 2, (box.y0 + box.y1) / 2)
            if any(c in r for r in later):
                hidden += not chr(ch[0]).isspace()
                chars.append(None)
            else:
                chars.append((chr(ch[0]), box))
        # split into runs of visible characters (a mask may cut a span in two)
        run: list = []
        for item in chars + [None]:
            if item is not None:
                run.append(item)
                continue
            text = "".join(c for c, _ in run)
            if text.strip():
                first = next(b for c, b in run if not c.isspace())
                spans.append((s["bbox"][1], round(first.x0), text.strip()))
            run = []
    return spans, hidden


def page_rows(page: pymupdf.Page, page_no: int, tol: float = 3.0, stats: dict | None = None) -> list[Row]:
    """Group visible text spans that share a baseline into rows of (x, text) cells."""
    spans, hidden = visible_spans(page)
    if stats is not None and hidden:
        stats.setdefault("hidden_chars", {})[page_no] = hidden
    spans.sort()
    rows: list[Row] = []
    for y, x, t in spans:
        if rows and abs(rows[-1].y - y) <= tol:
            rows[-1].cells.append((x, t))
        else:
            rows.append(Row(page_no, y, [(x, t)]))
    for r in rows:
        r.cells.sort()
    return rows


# --------------------------------------------------------------------------- result container

@dataclass
class Statement:
    source: str
    info: dict = field(default_factory=dict)
    account_summary: dict = field(default_factory=dict)
    sales_trend: list = field(default_factory=list)
    daily: list = field(default_factory=list)
    card_type: list = field(default_factory=list)
    location: list = field(default_factory=list)
    paid_by_others: list = field(default_factory=list)
    disputes: list = field(default_factory=list)
    fee_summary: list = field(default_factory=list)
    fees: list = field(default_factory=list)
    fee_category_totals: list = field(default_factory=list)
    interchange: list = field(default_factory=list)
    interchange_totals: list = field(default_factory=list)
    amounts_funded: list = field(default_factory=list)
    totals: dict = field(default_factory=dict)       # printed "Total" rows per section
    unparsed: list = field(default_factory=list)
    checks: list = field(default_factory=list)

    TABLES = ["daily", "card_type", "location", "paid_by_others", "disputes",
              "fee_summary", "fees", "fee_category_totals", "interchange",
              "interchange_totals", "amounts_funded", "sales_trend"]


# --------------------------------------------------------------------------- page 1 header

def parse_header(page: pymupdf.Page, rows: list[Row], st: Statement) -> None:
    text = "\n".join(" ".join(t for _, t in r.cells) for r in rows)   # visible text only
    if m := re.search(r"PERIOD:\s*(\d{2}/\d{2}/\d{4})\s*-\s*(\d{2}/\d{2}/\d{4})", text):
        st.info["period_start"], st.info["period_end"] = iso(m[1]), iso(m[2])
    if m := re.search(r"Corporate Number:\s*(\d+)", text):
        st.info["corporate_number"] = m[1]
    if m := re.search(r"Location Included:\s*(\d+)", text):
        st.info["locations_included"] = int(m[1])
    if m := re.search(r"Page \d+ of (\d+)", text):
        st.info["pages_printed"] = int(m[1])
    if m := re.search(r"Customer Service\s*(\S+)\s*([\d-]{10,})", text):
        st.info["customer_service_url"], st.info["customer_service_phone"] = m[1], m[2]
    elif (u := re.search(r"https?://\S+", text)) and (ph := re.search(r"\b1-\d{3}-\d{3}-\d{4}\b", text)):
        st.info["customer_service_url"], st.info["customer_service_phone"] = u[0], ph[0]

    # Merchant name/address block: left column between the statement title and the summary graphic
    addr = [t for r in rows if 125 < r.y < 200 for x, t in r.cells if x < 200]
    if addr:
        st.info["merchant_name"] = addr[0]
        st.info["merchant_address_lines"] = " | ".join(addr[1:])

    spans = [(r.y, x, t) for r in rows for x, t in r.cells]

    def value_below(label: str, pred=is_money, max_dy=60, max_dx=60):
        for y, x, t in spans:
            if t == label:
                cands = [(vy - y, abs(vx - x), vt) for vy, vx, vt in spans
                         if pred(vt) and 0 < vy - y <= max_dy and abs(vx - x) <= max_dx]
                if cands:
                    return min(cands)[2]
        return None

    for label, key in [("Amount Submitted", "amount_submitted"), ("Paid by Others", "paid_by_others"),
                       ("Disputes", "disputes"), ("Adjustments", "adjustments"), ("Fees", "fees"),
                       ("Amount Processed", "amount_processed")]:
        v = value_below(label)
        if v is not None:
            st.account_summary[key] = money(v)

    num = lambda s: bool(re.match(r"^[\d,]+(\.\d+)?$", s))
    for label, key in [("New Customers", "new_customers"), ("Avg. Ticket Size", "avg_ticket_size"),
                       ("Repeat Customers", "repeat_customers")]:
        v = value_below(label, pred=num, max_dy=20)
        if v is not None:
            st.info[key] = Decimal(v.replace(",", "")) if "." in v else to_int(v)

    # "Jul 25 vs. Jul 26" / "178.27%" / "year over year" are stacked; pair them by position
    def nearest(y, x, pred, max_dy=20):
        c = [(abs(vy - y) + abs(vx - x), vt) for vy, vx, vt in spans if pred(vt) and 0 < vy - y <= max_dy]
        return min(c)[1] if c else None
    for y, x, t in spans:
        if re.match(r"^[A-Z][a-z]{2} \d{2} vs\. [A-Z][a-z]{2} \d{2}$", t):
            p = nearest(y, x, lambda v: bool(PCT_RE.match(v)))
            kind = nearest(y, x, lambda v: v in ("year over year", "month over month"))
            if p and kind:
                st.info[kind.replace(" ", "_") + "_change"] = f"{t}: {p}"

    for y, x, t in spans:
        if m := re.match(r"^([A-Z][a-z]{2}) (\d{2}) - (-?\$[\d,]+\.\d{2})$", t):
            st.sales_trend.append({"month": datetime.strptime(f"{m[1]} {m[2]}", "%b %y").strftime("%Y-%m"),
                                   "amount_submitted": money(m[3]), "_x": x})
    st.sales_trend.sort(key=lambda d: d["month"])
    for d in st.sales_trend:
        d.pop("_x")


# --------------------------------------------------------------------------- section row parsers
# Each returns True if it consumed the row, False to send it to the "unparsed" log.

HEADER_WORDS = {"Date Submitted", "Location", "Card Type", "Category", "Type", "Product Description",
                "Items", "Total", "Date", "Submitted", "Amount", "Paid by Others", "Charges", "Program",
                "Sales", "% of", "Total Credits", "Service", "Interchange Charges", "Fees"}


def looks_like_header(r: Row) -> bool:
    return not any(is_money(t) or DATE_RE.match(t) or LOC_RE.match(t) for t in r.texts)


def p_daily(r: Row, st: Statement, ctx: dict) -> bool:
    t = r.texts
    if len(t) == 7 and all(is_money(v) for v in t[1:]):
        rec = dict(zip(["amount_submitted", "paid_by_others", "disputes", "adjustments", "fees",
                        "amount_processed"], map(money, t[1:])))
        if t[0] == "Total":
            st.totals["daily"] = rec
        else:
            st.daily.append({"date": iso(t[0]) if DATE_RE.match(t[0]) else t[0], **rec, "page": r.page})
        return True
    return False


def p_card_type(r: Row, st: Statement, ctx: dict) -> bool:
    t = r.texts
    # A long card name wraps around its values row: "AMERICAN" (line above) / values / "EXPRESS" (line below)
    if len(t) == 1 and t[0].isupper() and r.first_x < 60:
        last = st.card_type[-1] if st.card_type else None
        if last and ctx.get("wrapped") and r.page == last["page"] and r.y - ctx["last_y"] < 11:
            last["card_type"] += " " + t[0]
            ctx["wrapped"] = False
        else:
            ctx["pending"] = t[0]                     # first half of a wrapped name (or a chart legend label)
        return True
    nums = [v for v in t if is_money(v) or INT_RE.match(v)]
    names = [v for v in t if v not in nums]
    wrapped = not names and "pending" in ctx
    name = " ".join([ctx.pop("pending")] if wrapped else names).strip()
    if name == "Total" and len(nums) == 6:
        st.totals["card_type"] = dict(zip(["gross_items", "gross_amount", "refund_items", "refund_amount",
                                           "net_items", "net_amount"],
                                          [to_int(nums[0]), money(nums[1]), to_int(nums[2]), money(nums[3]),
                                           to_int(nums[4]), money(nums[5])]))
        return True
    if len(nums) == 7 and name:
        st.card_type.append({"card_type": name, "average_ticket": money(nums[0]),
                             "gross_items": to_int(nums[1]), "gross_amount": money(nums[2]),
                             "refund_items": to_int(nums[3]), "refund_amount": money(nums[4]),
                             "net_items": to_int(nums[5]), "net_amount": money(nums[6]), "page": r.page})
        ctx["wrapped"], ctx["last_y"] = wrapped, r.y
        return True
    if not nums or all(PCT_RE.match(v) for v in t):   # pie-chart labels and column headers
        ctx.pop("pending", None)
        return True
    return False


def p_location(r: Row, st: Statement, ctx: dict) -> bool:
    t = r.texts
    if len(t) == 7 and all(is_money(v) for v in t[1:]):
        rec = dict(zip(["amount_submitted", "paid_by_others", "disputes", "adjustments", "fees",
                        "amount_processed"], map(money, t[1:])))
        if t[0] == "Total":
            st.totals["location"] = rec
        else:
            st.location.append({"location": t[0], **rec, "page": r.page})
        return True
    return False


def p_paid_by_others(r: Row, st: Statement, ctx: dict) -> bool:
    t = r.texts
    if t[0] == "Total" and len(t) == 2 and is_money(t[1]):
        st.totals["paid_by_others"] = money(t[1])
        return True
    if len(t) == 3 and is_money(t[2]):
        st.paid_by_others.append({"location": t[0], "description": t[1], "amount": money(t[2]), "page": r.page})
        return True
    return False


def p_disputes(r: Row, st: Statement, ctx: dict) -> bool:
    t = r.texts
    if t[0] == "Total" and len(t) == 2 and is_money(t[1]):
        st.totals["disputes"] = money(t[1])
        return True
    if len(t) == 6 and DATE_RE.match(t[1]) and is_money(t[5]):
        st.disputes.append({"location": t[0], "date": iso(t[1]), "reference_no": t[2], "description": t[3],
                            "card_last4": t[4], "amount": money(t[5]), "page": r.page})
        return True
    # credits/debits banner above the table
    if all(is_money(v) or v in "-+=" or v.startswith("Total ") for v in t):
        ctx.setdefault("banner", []).extend(money(v) for v in t if is_money(v))
        return True
    return False


def p_fee_summary(r: Row, st: Statement, ctx: dict) -> bool:
    t = r.texts
    if t[0] == "Type" and len(t) == 8:
        ctx["brands"] = t[1:]
        return True
    if "brands" in ctx and len(t) == 8 and all(is_money(v) for v in t[1:]):
        vals = dict(zip(ctx["brands"], map(money, t[1:])))
        if t[0] == "Total":
            st.totals["fee_summary"] = vals
        else:
            st.fee_summary.append({"fee_type": t[0], **vals})
        return True
    # donut chart above the table: headline labels, numbers & percentages
    if "brands" not in ctx or looks_like_header(r):
        return True
    return False


FEE_DETAIL_PATTERNS = [
    (re.compile(r"([\d,]+) TRANS(?:ACTIONS)? AT \$?([\d.]+)"), ("item_count", "rate_per_item")),
    (re.compile(r"\$([\d,.]+) AT ?([\d.]+)"), ("basis_amount", "rate_pct")),
    (re.compile(r"([\d.]+) (?:DISC RATE )?TIMES \$([\d,.]+)"), ("rate_pct", "basis_amount")),
    (re.compile(r"TIMES \$([\d,.]+)"), ("basis_amount",)),
    (re.compile(r"(\d+) TRANS TOTALING \$([\d,.]+)"), ("item_count", "basis_amount")),
    (re.compile(r"([\d,]+) KILOBYTES AT ([\d.]+)"), ("item_count", "rate_per_item")),
]


def fee_details(desc: str) -> dict:
    out = {}
    for rx, names in FEE_DETAIL_PATTERNS:
        if m := rx.search(desc):
            for n, v in zip(names, m.groups()):
                out.setdefault(n, Decimal(v.replace(",", "")))
    return out


def p_fees(r: Row, st: Statement, ctx: dict) -> bool:
    t = r.texts
    if t[:2] == ["Category", "Product"]:
        ctx["cols"] = {name: x for x, name in r.cells}
        return True
    if t[0] == "Total" and len(t) == 2 and is_money(t[1]):
        st.fee_category_totals.append({"category": ctx.get("category"), "printed_total": money(t[1]),
                                       "page": r.page})
        return True
    cols = ctx.get("cols")
    if not cols or not is_money(t[-1]):
        return False
    x_prod, x_desc, x_type = cols["Product"] - 20, cols["Description"] - 20, cols["Type"] - 10
    rec = {"category": None, "product": None, "description": [], "fee_type": None}
    for x, v in r.cells[:-1]:
        if x < x_prod:
            rec["category"] = v
        elif x < x_desc:
            rec["product"] = v
        elif x < x_type:
            rec["description"].append(v)
        else:
            rec["fee_type"] = v
    # Category / Product are only printed on the first row of their group: carry them down
    if rec["category"]:
        ctx["category"], ctx["product"] = rec["category"], None
    if rec["product"]:
        ctx["product"] = rec["product"]
    desc = " ".join(rec["description"])
    st.fees.append({"category": ctx.get("category"), "product": ctx.get("product"), "description": desc,
                    "fee_type": rec["fee_type"], "amount": money(t[-1]), **fee_details(desc), "page": r.page})
    return True


IC_COLS = ["sales_total", "pct_of_sales", "transactions", "pct_of_transactions", "rate", "per_item_fee",
           "total_charges"]


def p_interchange(r: Row, st: Statement, ctx: dict) -> bool:
    t = r.texts
    if len(t) == 1 and t[0].isupper() and r.first_x < 60:
        ctx["brand"] = t[0]
        return True
    if len(t) == 8 and is_money(t[1]) and is_money(t[-1]):
        st.interchange.append({"card_brand": ctx.get("brand"), "product_description": t[0],
                               "sales_total": money(t[1]), "pct_of_sales": pct(t[2]),
                               "transactions": to_int(t[3]), "pct_of_transactions": pct(t[4]),
                               "rate": Decimal(t[5]), "per_item_fee": money(t[6]),
                               "total_charges": money(t[7]), "page": r.page})
        return True
    if t[0].endswith("Total") and len(t) == 4:
        rec = {"sales_total": money(t[1]), "transactions": to_int(t[2]), "total_charges": money(t[3])}
        if t[0] == "Total":
            st.totals["interchange"] = rec
        else:
            st.interchange_totals.append({"card_brand": t[0][:-len(" Total")], **rec, "page": r.page})
        return True
    return looks_like_header(r)


def p_amounts_funded(r: Row, st: Statement, ctx: dict) -> bool:
    t = r.texts
    if len(t) == 8 and DATE_RE.match(t[0]) and DATE_RE.match(t[6]):
        st.amounts_funded.append({"date_submitted": iso(t[0]),
                                  **dict(zip(["amount_submitted", "paid_by_others", "disputes", "adjustments",
                                              "fees"], map(money, t[1:6]))),
                                  "date_funded": iso(t[6]), "amount_funded": money(t[7]), "page": r.page})
        return True
    if t[0] == "Total" and len(t) == 7:
        st.totals["amounts_funded"] = dict(zip(["amount_submitted", "paid_by_others", "disputes", "adjustments",
                                                "fees", "amount_funded"], map(money, t[1:])))
        return True
    return looks_like_header(r) or t[0].startswith(("This section", "your account", "fees reported"))


PARSERS = {"daily": p_daily, "card_type": p_card_type, "location": p_location,
           "paid_by_others": p_paid_by_others, "disputes": p_disputes, "fee_summary": p_fee_summary,
           "fees": p_fees, "interchange": p_interchange, "amounts_funded": p_amounts_funded}


# --------------------------------------------------------------------------- driver

def is_page_chrome(r: Row) -> bool:
    return r.y < 75 or any(re.match(r"^Page \d+ of \d+$", t) or t.startswith("PERIOD:") for t in r.texts)


def parse(pdf_path: str | Path) -> Statement:
    doc = pymupdf.open(pdf_path)
    st = Statement(source=str(pdf_path))
    section, ctx, ctx_stats = None, {}, {}
    for pno, page in enumerate(doc, start=1):
        rows = page_rows(page, pno, stats=ctx_stats)
        if pno == 1:
            parse_header(page, rows, st)
        for r in rows:
            if is_page_chrome(r):
                continue
            title = r.texts[0]
            if TITLE_X[0] <= r.first_x <= TITLE_X[1] and title in SECTION_TITLES:
                new = SECTION_TITLES[title]
                if new != section:                 # same title repeated on a continuation page keeps context
                    section, ctx = new, {}
                continue
            if section is None:
                continue
            if not PARSERS[section](r, st, ctx) and not looks_like_header(r):
                st.unparsed.append({"page": r.page, "section": section, "text": " | ".join(r.texts)})
    st.info["pdf_pages"] = len(doc)
    validate(st)
    if hidden := ctx_stats.get("hidden_chars"):
        st.checks.append({"check": "Text hidden under shapes (masked / redacted areas)", "expected": None,
                          "actual": sum(hidden.values()), "difference": None, "status": "REVIEW",
                          "note": f"{sum(hidden.values())} characters on pages {sorted(hidden)} are covered by "
                                  f"filled boxes and were excluded. The masked text is still present in the PDF "
                                  f"text layer, so other tools can still read it."})
    if st.info.get("pages_printed") not in (None, len(doc)):
        blank = [i + 1 for i, pg in enumerate(doc) if not pg.get_text().strip()]
        st.checks.append({"check": "Page count", "expected": st.info["pages_printed"], "actual": len(doc),
                          "difference": None, "status": "INFO",
                          "note": f"statement says 'of {st.info['pages_printed']}', PDF has {len(doc)} pages; "
                                  f"blank pages: {blank}"})
    return st


# --------------------------------------------------------------------------- reconciliation

def validate(st: Statement) -> None:
    def check(name, expected, actual, note=""):
        if expected is None or actual is None:
            st.checks.append({"check": name, "expected": expected, "actual": actual, "difference": None,
                              "status": "MISSING", "note": note or "value not found in PDF"})
            return
        diff = Decimal(actual) - Decimal(expected)
        st.checks.append({"check": name, "expected": expected, "actual": actual, "difference": diff,
                          "status": "OK" if abs(diff) <= TOLERANCE else "MISMATCH", "note": note})

    def flag(name, note, status="REVIEW"):
        st.checks.append({"check": name, "expected": None, "actual": None, "difference": None,
                          "status": status, "note": note})

    S = lambda rows, k, **flt: sum((r[k] for r in rows if all(r.get(a) == b for a, b in flt.items())), Decimal(0))
    a, T = st.account_summary, st.totals

    # Account summary arithmetic
    if len(a) == 6:
        check("Account summary: submitted + paid by others + disputes + adjustments + fees = processed",
              a["amount_processed"],
              a["amount_submitted"] + a["paid_by_others"] + a["disputes"] + a["adjustments"] + a["fees"])

    # Summary by day
    days = [d for d in st.daily]
    for k in ["amount_submitted", "paid_by_others", "disputes", "adjustments", "fees", "amount_processed"]:
        check(f"Summary by Day: sum of rows = printed Total ({k})", T.get("daily", {}).get(k), S(days, k))
        check(f"Summary by Day Total = Account Summary ({k})", a.get(k), T.get("daily", {}).get(k))
    bad = [d["date"] for d in days if abs(d["amount_submitted"] + d["paid_by_others"] + d["disputes"]
                                          + d["adjustments"] + d["fees"] - d["amount_processed"]) > TOLERANCE]
    check("Summary by Day: every row adds across (count of rows that don't)", 0, len(bad), ", ".join(bad))

    # Card type
    ct = T.get("card_type", {})
    for k in ["gross_items", "gross_amount", "refund_items", "refund_amount", "net_items", "net_amount"]:
        check(f"Card Type: sum of rows = printed Total ({k})", ct.get(k), S(st.card_type, k))
    check("Card Type net amount = Account Summary amount submitted", a.get("amount_submitted"), ct.get("net_amount"))
    bad = [c["card_type"] for c in st.card_type
           if c["gross_amount"] + c["refund_amount"] != c["net_amount"] or c["gross_items"] + c["refund_items"] != c["net_items"]]
    check("Card Type: gross + refunds = net on every row (rows that don't)", 0, len(bad), ", ".join(bad))
    amex = next((c for c in st.card_type if "AMERICAN" in c["card_type"]), None)
    if amex:
        check("AMEX net volume = Paid by Others (AMEX is funded directly by American Express)",
              -a.get("paid_by_others", 0), amex["net_amount"])

    # Location
    for k in ["amount_submitted", "paid_by_others", "disputes", "adjustments", "fees", "amount_processed"]:
        check(f"Summary by Location: sum of rows = printed Total ({k})", T.get("location", {}).get(k), S(st.location, k))
    if "locations_included" in st.info:
        check("Number of locations listed = 'Location Included' on page 1", st.info["locations_included"], len(st.location))

    # Paid by others & disputes
    check("Paid by Others: sum of detail = printed Total", T.get("paid_by_others"), S(st.paid_by_others, "amount"))
    check("Paid by Others Total = Account Summary", a.get("paid_by_others"), T.get("paid_by_others"))
    loc_pbo = {l["location"]: l["paid_by_others"] for l in st.location}
    bad = [p["location"] for p in st.paid_by_others if loc_pbo.get(p["location"]) != p["amount"]]
    check("Paid by Others detail matches Summary by Location per location (mismatching)", 0, len(bad), ", ".join(bad))
    check("Disputes: sum of detail = printed Total", T.get("disputes"), S(st.disputes, "amount"))
    check("Disputes Total = Account Summary", a.get("disputes"), T.get("disputes"))

    # Fees
    check("Fee detail: sum of all fee lines = Account Summary fees", a.get("fees"), S(st.fees, "amount"))
    fs = {r["fee_type"]: r for r in st.fee_summary}
    for ftype, row in fs.items():
        check(f"Fee detail lines of type '{ftype}' = Fee Summary '{ftype}' total", row.get("Total"),
              S(st.fees, "amount", fee_type=ftype))
    if "fee_summary" in T:
        check("Fee Summary grand total = Account Summary fees", a.get("fees"), T["fee_summary"].get("Total"))
    # printed category subtotals (a category may continue across pages; subtotal follows its last row)
    running, i = Decimal(0), 0
    for f in st.fees:
        f["_seq"] = i; i += 1
    for tot in st.fee_category_totals:
        check(f"Fee category subtotal '{tot['category']}' (page {tot['page']})", tot["printed_total"],
              S(st.fees, "amount", category=tot["category"]))
    for f in st.fees:
        f.pop("_seq", None)
    missing_type = [f["description"] for f in st.fees if not f["fee_type"]]
    check("Fee lines without a fee type", 0, len(missing_type), "; ".join(missing_type[:5]))

    # Interchange
    for brand in {r["card_brand"] for r in st.interchange}:
        pt = next((t for t in st.interchange_totals if t["card_brand"] == brand), None)
        rows = [r for r in st.interchange if r["card_brand"] == brand]
        for k in ["sales_total", "transactions", "total_charges"]:
            check(f"Interchange {brand}: sum of rows = printed brand total ({k})", pt and pt[k], S(rows, k))
    ic = T.get("interchange", {})
    for k in ["sales_total", "transactions", "total_charges"]:
        check(f"Interchange: sum of brand totals = printed grand Total ({k})", ic.get(k), S(st.interchange_totals, k))
    check("Interchange sales total = Account Summary amount submitted", a.get("amount_submitted"), ic.get("sales_total"))
    check("Interchange transaction count = Card Type net items", ct.get("net_items"), ic.get("transactions"))
    by_brand = {c["card_type"].replace(" ", ""): c for c in st.card_type}
    for t in st.interchange_totals:
        c = by_brand.get(t["card_brand"].replace(" ", ""))
        if c:
            check(f"Interchange {t['card_brand']} sales = Card Type net amount", c["net_amount"], t["sales_total"])
    # Recompute each interchange line from its printed rate; small differences are normal (refund
    # netting, per-item rounding), larger ones are worth a look by the analyst.
    for r in st.interchange:
        calc = -(r["sales_total"] * r["rate"] + r["transactions"] * r["per_item_fee"])
        r["recalculated_charges"] = calc.quantize(Decimal("0.01"))
        r["recalc_difference"] = (r["total_charges"] - r["recalculated_charges"])
    bad = [r for r in st.interchange if r["total_charges"] and abs(r["recalc_difference"]) > Decimal("1")]
    if bad:
        flag("Interchange lines that don't recompute from printed rate (> $1)",
             "; ".join(f"{r['card_brand']} {r['product_description']} (p{r['page']}): printed {r['total_charges']}, "
                       f"sales x rate + items x per-item = {r['recalculated_charges']}" for r in bad))
    if "Interchange Charges" in fs and ic:
        # Fee Summary "Interchange Charges" = interchange table + debit network interchange (the table shows
        # $0.00 for debit) + card-brand assessments that are typed as Interchange Charges in the fee detail.
        ic_fees = [f for f in st.fees if f["fee_type"] == "Interchange Charges"]
        debit = S(ic_fees, "amount", product="Debit")
        names = {r["product_description"] for r in st.interchange}
        other = sum((f["amount"] for f in ic_fees if f["product"] != "Debit" and f["description"] not in names),
                    Decimal(0))
        check("Fee Summary Interchange = interchange table + debit network IC + assessments/network fees",
              fs["Interchange Charges"]["Total"], ic["total_charges"] + debit + other,
              f"interchange table {ic['total_charges']} + debit network interchange {debit} "
              f"+ brand assessments/network fees {other}")

    # Amounts funded
    af = T.get("amounts_funded", {})
    for k in ["amount_submitted", "paid_by_others", "disputes", "adjustments", "fees", "amount_funded"]:
        check(f"Amounts Funded: sum of rows = printed Total ({k})", af.get(k), S(st.amounts_funded, k))
    bad = [r["date_submitted"] for r in st.amounts_funded
           if abs(r["amount_submitted"] + r["paid_by_others"] + r["disputes"] + r["adjustments"] + r["fees"]
                  - r["amount_funded"]) > TOLERANCE]
    check("Amounts Funded: every row adds across (rows that don't)", 0, len(bad), ", ".join(bad))
    check("Amounts Funded submitted = Account Summary submitted", a.get("amount_submitted"), af.get("amount_submitted"))
    if af and a:
        flag("Funded vs processed", f"Amount funded {af.get('amount_funded')} vs amount processed "
             f"{a.get('amount_processed')}: fees collected in this period ({af.get('fees')}) are the prior "
             f"statement's fees, while this statement's fees ({a.get('fees')}) are charged as a month-end "
             f"charge and collected next period.", "INFO")
        pos_pbo = [r for r in st.amounts_funded if r["paid_by_others"] > 0]
        if pos_pbo:
            flag("Positive 'Paid by Others' in Amounts Funded",
                 "; ".join(f"{r['date_submitted']}: {r['paid_by_others']}" for r in pos_pbo)
                 + " (usually negative; likely an AMEX refund netted into funding)")

    # Period / trend
    ps, pe = st.info.get("period_start"), st.info.get("period_end")
    out = [d["date"] for d in st.daily if re.match(r"\d{4}-", d["date"]) and ps and not (ps <= d["date"] <= pe)]
    if out:
        flag("Summary by Day dates outside statement period", ", ".join(out)
             + " (batches submitted before the period but posted inside it)")
    if st.sales_trend and pe:
        last = st.sales_trend[-1]
        if last["month"] == pe[:7]:
            check("13-month trend current month = amount submitted", a.get("amount_submitted"), last["amount_submitted"])
    check("Fee category subtotals add up to Account Summary fees", a.get("fees"),
          S(st.fee_category_totals, "printed_total"))
    if st.unparsed:
        flag("Unparsed rows", f"{len(st.unparsed)} row(s) could not be mapped; see Unparsed sheet", "MISMATCH")


# --------------------------------------------------------------------------- output

def _plain(v):
    if isinstance(v, Decimal):
        return float(v)
    return v


def write_excel(st: Statement, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1F3864")
    fills = {"OK": "C6EFCE", "MISMATCH": "FFC7CE", "MISSING": "FFEB9C", "REVIEW": "FFEB9C", "INFO": "DDEBF7"}

    def sheet(name, rows, money_cols=(), pct_cols=()):
        ws = wb.create_sheet(name)
        if not rows:
            ws.append(["(no rows)"])
            return ws
        cols = list(dict.fromkeys(k for r in rows for k in r))
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append([_plain(r.get(c)) for c in cols])
        for i, c in enumerate(cols, start=1):
            letter = get_column_letter(i)
            width = max(len(str(c)), *(len(str(_plain(r.get(c, "")))) for r in rows))
            ws.column_dimensions[letter].width = min(max(10, width + 2), 70)
            fmt = ("#,##0.00;[Red]-#,##0.00" if c in money_cols else "0.00%" if c in pct_cols else None)
            if fmt:
                for cell in ws[letter][1:]:
                    cell.number_format = fmt
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    M = {"amount_submitted", "paid_by_others", "disputes", "adjustments", "fees", "amount_processed",
         "amount", "average_ticket", "gross_amount", "refund_amount", "net_amount", "sales_total",
         "per_item_fee", "total_charges", "amount_funded", "printed_total", "basis_amount", "expected",
         "actual", "difference", "Visa", "Mastercard", "Amex", "Discover", "Debit", "Others", "Total"}

    ws = sheet("Validation", st.checks, M)
    for row in ws.iter_rows(min_row=2):
        status = row[4].value
        if status in fills:
            for c in row:
                c.fill = PatternFill("solid", fgColor=fills[status])
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["F"].width = 90

    info = [{"field": k, "value": _plain(v)} for k, v in st.info.items()]
    info += [{"field": f"account_summary.{k}", "value": _plain(v)} for k, v in st.account_summary.items()]
    info.append({"field": "source_file", "value": st.source})
    sheet("Statement_Info", info)
    sheet("Daily_Summary", st.daily, M)
    sheet("Card_Type", st.card_type, M)
    sheet("Location", st.location, M)
    sheet("Paid_by_Others", st.paid_by_others, M)
    sheet("Disputes", st.disputes, M)
    sheet("Fee_Summary", st.fee_summary, M)
    sheet("Fee_Detail", st.fees, M | {"rate_per_item"}, {"rate_pct"})
    sheet("Fee_Category_Totals", st.fee_category_totals, M)
    sheet("Interchange", st.interchange, M, {"pct_of_sales", "pct_of_transactions", "rate"})
    sheet("Interchange_Brand_Totals", st.interchange_totals, M)
    sheet("Amounts_Funded", st.amounts_funded, M)
    sheet("Sales_Trend_13M", st.sales_trend, M)
    sheet("Unparsed", st.unparsed)
    wb.move_sheet("Validation", offset=1)  # Statement_Info first, then the checks, then the data
    wb.save(path)


def write_csv(st: Statement, folder: Path) -> None:
    import csv
    folder.mkdir(parents=True, exist_ok=True)
    for name in Statement.TABLES + ["checks", "unparsed"]:
        rows = getattr(st, name)
        if not rows:
            continue
        cols = list(dict.fromkeys(k for r in rows for k in r))
        with open(folder / f"{name}.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, cols)
            w.writeheader()
            w.writerows({k: _plain(v) for k, v in r.items()} for r in rows)


def write_json(st: Statement, path: Path) -> None:
    data = {k: getattr(st, k) for k in ["info", "account_summary", "totals", *Statement.TABLES, "checks", "unparsed"]}
    data["source"] = st.source
    path.write_text(json.dumps(data, indent=2, default=lambda v: str(v) if isinstance(v, Decimal) else v))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", help="statement PDF, or a folder of PDFs")
    ap.add_argument("-o", "--output", help="Excel output path (single PDF) or folder (batch)")
    ap.add_argument("--csv", help="also write one CSV per table into this folder")
    ap.add_argument("--json", help="also write a JSON file")
    args = ap.parse_args(argv)

    src = Path(args.input)
    pdfs = sorted(src.glob("*.pdf")) if src.is_dir() else [src]
    worst = 0
    for pdf in pdfs:
        st = parse(pdf)
        if len(pdfs) == 1 and args.output:
            out = Path(args.output)
        else:
            out = (Path(args.output) if args.output else pdf.parent) / f"{pdf.stem}_parsed.xlsx"
            out.parent.mkdir(parents=True, exist_ok=True)
        write_excel(st, out)
        if args.csv:
            write_csv(st, Path(args.csv) / pdf.stem if len(pdfs) > 1 else Path(args.csv))
        if args.json:
            write_json(st, Path(args.json) if len(pdfs) == 1 else Path(args.json) / f"{pdf.stem}.json")

        counts = {s: sum(c["status"] == s for c in st.checks) for s in ["OK", "MISMATCH", "MISSING", "REVIEW", "INFO"]}
        print(f"{pdf.name}: {len(st.daily)} days, {len(st.location)} locations, {len(st.fees)} fee lines, "
              f"{len(st.interchange)} interchange lines, {len(st.amounts_funded)} funding rows -> {out}")
        print("   checks: " + ", ".join(f"{k} {v}" for k, v in counts.items() if v))
        for c in st.checks:
            if c["status"] in ("MISMATCH", "MISSING"):
                print(f"   ! {c['status']}: {c['check']} (expected {c['expected']}, got {c['actual']}) {c['note']}")
                worst = 1
    return worst


if __name__ == "__main__":
    sys.exit(main())
